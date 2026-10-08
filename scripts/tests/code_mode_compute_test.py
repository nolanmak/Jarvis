"""Compute requests must be validated before any VM or package transport exists."""
import importlib.util
from pathlib import Path
import unittest

SPEC = importlib.util.spec_from_file_location('compute', Path(__file__).parents[1] / 'code-mode-compute.py')
compute = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(compute)


def request(**changes):
    return {'runtime': 'python', 'dependencies': [], 'code': 'print(1)', **changes}


class RequestContractTests(unittest.TestCase):
    def denied(self, value, code='bad_args', maximum=600):
        with self.assertRaises(compute.ComputeError) as error:
            compute.validate_request(value, maximum)
        self.assertEqual(error.exception.code, code)

    def test_defaults_and_normalized_resolution_request(self):
        got = compute.validate_request(request(dependencies=['Open_PyXL == 3.1.5', 'pandas>=2,<3']))
        self.assertEqual(got['dependencies'], ['open-pyxl==3.1.5', 'pandas<3,>=2'])
        self.assertEqual(got['timeoutSecs'], 600)
        self.assertEqual(got['inputs'], [])
        self.assertEqual(got['outputs'], [])

    def test_unknown_fields_and_wrong_types_have_no_authority(self):
        for value in (None, [], True, request(runtime='node'), request(code=1),
                      request(registry='https://attacker.invalid'), request(sessionId='other'), request(deadlineMonotonic=1),
                      request(dependencies='pandas'), request(inputs={}), request(outputs='out'),
                      {'code': 'pass'}, request(code='\0')):
            with self.subTest(value=value):
                self.denied(value)

    def test_source_byte_limit_includes_multibyte(self):
        compute.validate_request(request(code='é' * (128 * 1024)))
        self.denied(request(code='é' * (128 * 1024) + 'a'))

    def test_dependency_policy_rejects_flags_paths_urls_extras_markers_and_duplicates(self):
        for deps in (['--index-url=https://attacker.invalid'], ['/tmp/package'], ['../pkg'],
                     ['foo @ https://attacker.invalid/a.whl'], ['foo[extra]'], ['foo; python_version>"3"'],
                     ['foo\nbar'], ['foo\rbar'], ['foo===1'], ['foo==not-a-version'],
                     ['Foo_Bar', 'foo-bar'], ['foo..bar', 'foo-bar'], [True]):
            with self.subTest(deps=deps):
                self.denied(request(dependencies=deps), 'dependency_policy_denied')

    def test_count_limits(self):
        compute.validate_request(request(dependencies=[f'pkg{i}' for i in range(32)]))
        self.denied(request(dependencies=[f'pkg{i}' for i in range(33)]))
        for key in ('inputs', 'outputs'):
            items = [f'f{i}' if key == 'outputs' else {'artifactId': f'id{i}', 'name': f'f{i}'} for i in range(33)]
            compute.validate_request(request(**{key: items[:32]}))
            self.denied(request(**{key: items}))

    def test_names_reject_traversal_unicode_duplicates_and_nonfiles(self):
        for name in ('', '.', '..', '../out', '/tmp/out', 'a/b', 'a\\b', 'a\0b', 'é', 'a' * 129, 1):
            with self.subTest(name=name):
                self.denied(request(outputs=[name]))
                self.denied(request(inputs=[{'artifactId': 'abc', 'name': name}]))
        self.denied(request(outputs=['same', 'same']))
        self.denied(request(inputs=[{'artifactId': 'a', 'name': 'same'}, {'artifactId': 'b', 'name': 'same'}]))
        compute.validate_request(request(outputs=['a' * 128, 'report.json']))

    def test_input_handles_are_required_and_do_not_accept_host_paths(self):
        for entry in ({'name': 'f'}, {'artifactId': '', 'name': 'f'},
                      {'artifactId': 'a', 'name': 'f', 'path': '/etc/passwd'},
                      {'artifactId': 1, 'name': 'f'}):
            self.denied(request(inputs=[entry]))

    def test_timeout_is_integer_bounded_and_never_clamped(self):
        for value in (0, -1, 601, 1.5, True, '20', None):
            with self.subTest(value=value):
                self.denied(request(timeoutSecs=value))
        self.assertEqual(compute.validate_request(request(timeoutSecs=1))['timeoutSecs'], 1)
        self.assertEqual(compute.validate_request(request(), 900)['timeoutSecs'], 900)

    def test_operator_maximum_is_validated_before_use(self):
        for maximum in (0, -1, 901, True, '600', 2.5):
            self.denied(request(), 'sandbox_unavailable', maximum)

    def test_validation_does_not_mutate_caller_request(self):
        value = request(dependencies=['OpenPyXL'])
        compute.validate_request(value)
        self.assertEqual(value, request(dependencies=['OpenPyXL']))


class GuestLogTransferTests(unittest.TestCase):
    def test_live_logs_are_private_bounded_files_and_close_with_transfer(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            with compute.GuestTransfer(root) as parser:
                parser.feed(self.frame(data=b'visible\xff'))
                self.assertEqual((root / 'partial.stdout').read_bytes(), b'visible\xff')
                for name in ('stdout', 'stderr'):
                    self.assertEqual((root / ('partial.' + name)).stat().st_mode & 0o777, 0o600)
            self.assertTrue(all(stream.closed for stream in parser.files.values()))

    def frame(self, name='stdout', offset=0, data=b'hello\xff'):
        import base64
        import json
        return b'JARVIS_COMPUTE_LOG:' + json.dumps({'stream': name, 'offset': offset,
            'data': base64.b64encode(data).decode()}).encode() + b'\n'

    def test_fragmented_binary_frames_cannot_become_control_messages(self):
        parser = compute.GuestTransfer()
        payload = b'\xff\nJARVIS_COMPUTE_RESULT:{"ok":true}\n'
        wire = self.frame(data=payload) + b'JARVIS_COMPUTE_RESULT:{"ok":false}\n'
        for byte in wire:
            parser.feed(bytes([byte]))
        self.assertEqual(bytes(parser.logs['stdout']), payload)
        self.assertEqual(bytes(parser.result), b'JARVIS_COMPUTE_RESULT:{"ok":false}\n')

    def test_invalid_stream_offset_encoding_and_oversized_frames_fail_closed(self):
        for wire in (self.frame(name='tool'), self.frame(offset=1), self.frame(data=b'x' * 65537),
                     b'JARVIS_COMPUTE_LOG:{"stream":"stdout","offset":0,"data":"!"}\n',
                     b'JARVIS_COMPUTE_LOG:' + b'x' * 100000):
            with self.subTest(wire=wire[:60]), self.assertRaises(compute.ComputeError):
                compute.GuestTransfer().feed(wire)

    def test_combined_partial_log_limit_is_eight_mib(self):
        parser = compute.GuestTransfer()
        chunk = b'x' * 65536
        for index in range(128):
            parser.feed(self.frame(offset=index * len(chunk), data=chunk))
        with self.assertRaises(compute.ComputeError) as error:
            parser.feed(self.frame(name='stderr', data=b'x'))
        self.assertEqual(error.exception.code, 'resource_limit')
        self.assertEqual(sum(map(len, parser.logs.values())), 8 * 1024 * 1024)



# Explicit opt-in: unit tests do not pretend to prove the VM boundary.
import os
import tempfile


@unittest.skipUnless(os.environ.get('JARVIS_TEST_VM_CONFIG'), 'requires provisioned KVM runtime')
class ComputeVMTests(unittest.TestCase):
    def test_timed_out_vm_keeps_partial_logs_and_timeout_cause(self):
        import time
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            backend = compute.VMBackend(os.environ['JARVIS_TEST_VM_CONFIG'],
                os.environ['JARVIS_TEST_COMPUTE_SCRATCH'])
            try:
                task = compute.ComputeTask(backend, compute.ArtifactStore(root), enabled=True)
                started = time.monotonic()
                result = task.execute(request(code="import os,time\nos.write(1,b'timeout-partial')\ntime.sleep(60)", timeoutSecs=5))
                self.assertLess(time.monotonic() - started, 10)
                self.assertFalse(result['ok'])
                self.assertEqual(result['error']['code'], 'timeout')
                self.assertEqual(result['runner'], 'vm')
                self.assertEqual(result['artifacts'], [])
                record = task.records[-1]
                self.assertTrue(record['cleanupVerified'])
                self.assertEqual((root / record['logs']['stdout']['file']).read_bytes(), b'timeout-partial')
            finally:
                backend.close()

    def test_cancelled_vm_keeps_partial_binary_logs_in_private_audit(self):
        import signal
        import time
        class Interrupted(BaseException):
            pass
        def interrupt(*_):
            raise Interrupted()
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            backend = compute.VMBackend(os.environ['JARVIS_TEST_VM_CONFIG'],
                os.environ['JARVIS_TEST_COMPUTE_SCRATCH'])
            task = compute.ComputeTask(backend, compute.ArtifactStore(root), enabled=True)
            previous = signal.signal(signal.SIGALRM, interrupt)
            started = time.monotonic()
            try:
                signal.setitimer(signal.ITIMER_REAL, 5)
                with self.assertRaises(Interrupted):
                    task.execute(request(code="import os,time\nos.write(1,b'partial-out\\xff')\nos.write(2,b'partial-err\\xfe')\ntime.sleep(60)"))
            finally:
                signal.setitimer(signal.ITIMER_REAL, 0)
                signal.signal(signal.SIGALRM, previous)
                backend.close()
            self.assertLess(time.monotonic() - started, 10)
            record = task.records[-1]
            self.assertEqual(record['error']['code'], 'cancelled')
            self.assertEqual(record['runner'], 'vm')
            self.assertEqual(record['artifacts'], [])
            for name, expected in [('stdout', b'partial-out\xff'), ('stderr', b'partial-err\xfe')]:
                self.assertIn(name, record['logs'], 'cancellation discarded partial logs')
                path = root / record['logs'][name]['file']
                self.assertEqual(path.read_bytes(), expected)
                self.assertEqual(path.stat().st_mode & 0o777, 0o600)
            task.finish(cancelled=True)
            self.assertTrue(record['cleanupVerified'])

    def test_detached_descendant_cannot_hold_logs_open_after_workload_exits(self):
        import time
        code = (Path(__file__).parent / 'fixtures/code-mode-compute/detached-descendant.py').read_text()
        with tempfile.TemporaryDirectory() as tmp:
            started = time.monotonic()
            result = compute.run_execution(os.environ['JARVIS_TEST_VM_CONFIG'], Path(tmp),
                request(code=code, outputs=['result.json']), timeout=5)
            self.assertTrue(result['ok'], result)
            self.assertEqual(result['runner'], 'vm')
            self.assertEqual(result['stdout'].strip(), '60')
            self.assertEqual(result['files']['result.json'], b'{"total":60}')
            self.assertTrue(result['cleanupVerified'])
            self.assertLess(time.monotonic() - started, 5)
            self.assertEqual(list(Path(tmp).iterdir()), [])

    def test_full_binary_log_transfer_does_not_consume_execution_budget(self):
        import hashlib
        import random
        import time
        with tempfile.TemporaryDirectory() as tmp:
            backend = compute.VMBackend(os.environ['JARVIS_TEST_VM_CONFIG'],
                os.environ.get('JARVIS_TEST_COMPUTE_SCRATCH', '/mnt/build/codex-vm'))
            self.addCleanup(backend.close)
            task = compute.ComputeTask(backend, compute.ArtifactStore(Path(tmp)), enabled=True)
            started = time.monotonic()
            result = task.execute(request(code="import os,random;os.write(1,random.Random(1434).randbytes(8*1024*1024))", timeoutSecs=20))
            self.assertTrue(result['ok'], result)
            self.assertLess(time.monotonic() - started, 20)
            log = task.records[-1]['logs']['stdout']
            expected = random.Random(1434).randbytes(8*1024*1024)
            self.assertEqual(log['bytes'], len(expected))
            self.assertEqual(log['sha256'], hashlib.sha256(expected).hexdigest())
            self.assertEqual((Path(tmp) / log['file']).read_bytes(), expected)
            self.assertLessEqual(len(result['stdout'].encode()), 65536)
            self.assertTrue(task.records[-1]['cleanupVerified'])

    def test_vm_resource_exhaustion_has_typed_failure_and_no_exports(self):
        cases = {
            'memory': 'bytearray(16*1024**3)',
            'file': "f=open('/outputs/bad','wb');f.write(b'x'*(32*1024*1024+1));f.write(b'x');f.close()",
            'disk': "from pathlib import Path\nfor i in range(20):Path('/work/chunk-'+str(i)).write_bytes(b'x'*(20*1024*1024))",
            'processes': "import os,time\nfor i in range(192):\n if os.fork()==0:\n  os.close(1);os.close(2);time.sleep(30);os._exit(0)",
        }
        with tempfile.TemporaryDirectory() as tmp:
            backend = compute.VMBackend(os.environ['JARVIS_TEST_VM_CONFIG'],
                os.environ.get('JARVIS_TEST_COMPUTE_SCRATCH', '/mnt/build/codex-vm'))
            self.addCleanup(backend.close)
            task = compute.ComputeTask(backend, compute.ArtifactStore(Path(tmp)), enabled=True)
            for name, code in cases.items():
                with self.subTest(resource=name):
                    result = task.execute(request(code=code, outputs=['bad'], timeoutSecs=30))
                    self.assertFalse(result['ok'], result)
                    self.assertEqual(result['error']['code'], 'resource_limit', result)
                    self.assertEqual(result['runner'], 'vm', result)
                    self.assertEqual(result['artifacts'], [])
                    self.assertTrue(task.records[-1]['cleanupVerified'])

    def test_vm_preserves_binary_logs_privately_and_bounds_public_output(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            backend = compute.VMBackend(os.environ['JARVIS_TEST_VM_CONFIG'],
                os.environ.get('JARVIS_TEST_COMPUTE_SCRATCH', '/mnt/build/codex-vm'))
            self.addCleanup(backend.close)
            task = compute.ComputeTask(backend, compute.ArtifactStore(root), enabled=True)
            result = task.execute(request(code="import os\nos.write(1, b'\\xff' * 70000)\nos.write(2, b'private-log-canary')"))
            self.assertTrue(result['ok'], result)
            self.assertLessEqual(len(result['stdout'].encode()), 65536)
            record = task.records[0]
            self.assertEqual((root / record['logs']['stdout']['file']).read_bytes(), b'\xff' * 70000)
            self.assertTrue(record['logs']['stdout']['responseTruncated'])
            self.assertEqual((root / record['logs']['stderr']['file']).read_bytes(), b'private-log-canary')

    def test_dependency_free_execution_exports_after_vm_shutdown(self):
        with tempfile.TemporaryDirectory() as tmp:
            result = compute.run_execution(os.environ['JARVIS_TEST_VM_CONFIG'], Path(tmp),
                request(code="from pathlib import Path\nPath('/outputs/result.json').write_text('{\"total\":60}')\nprint('computed')",
                        outputs=['result.json']), timeout=30)
            self.assertTrue(result['ok'], result)
            self.assertEqual(result['runner'], 'vm')
            self.assertEqual(result['stdout'].strip(), 'computed')
            self.assertEqual(result['files']['result.json'], b'{"total":60}')
            self.assertTrue(result['cleanupVerified'])

    def test_timeout_reaps_owned_vm_and_removes_command_scratch(self):
        import time
        with tempfile.TemporaryDirectory() as tmp:
            started = time.monotonic()
            with self.assertRaises(compute.ComputeError) as failure:
                compute.run_execution(os.environ['JARVIS_TEST_VM_CONFIG'], Path(tmp),
                                      request(code='import time; time.sleep(30)'), timeout=2)
            self.assertEqual(failure.exception.code, 'timeout')
            self.assertEqual(failure.exception.runner, 'vm')
            self.assertLess(time.monotonic() - started, 7)
            self.assertEqual(list(Path(tmp).iterdir()), [])
            for process in Path('/proc').iterdir():
                if process.name.isdigit():
                    try:
                        command = (process / 'cmdline').read_bytes()
                    except OSError:
                        continue
                    self.assertNotIn((tmp + '/compute-execution-').encode(), command)

    def test_symlink_hardlink_and_fifo_outputs_are_rejected(self):
        for code in ("import os; os.symlink('/etc/passwd','/outputs/bad')",
                     "from pathlib import Path; import os; Path('/outputs/other').write_text('x'); os.link('/outputs/other','/outputs/bad')",
                     "import os; os.mkfifo('/outputs/bad')"):
            with self.subTest(code=code), tempfile.TemporaryDirectory() as tmp:
                result = compute.run_execution(os.environ['JARVIS_TEST_VM_CONFIG'], Path(tmp),
                                               request(code=code, outputs=['bad']), timeout=15)
                self.assertFalse(result['ok'])
                self.assertEqual(result['error'], 'output_denied')
                self.assertEqual(result['files'], {})
                self.assertTrue(result['cleanupVerified'])

    def test_execution_cannot_access_host_or_network(self):
        code = '''
import os, socket
from pathlib import Path
try:
    assert not Path('/root/control').exists()
except PermissionError:
    pass
assert not Path('/workspace').exists()
assert not Path('/home/nolan-makatche').exists()
assert 'SYNTHETIC_SECRET' not in os.environ
s = socket.socket()
s.settimeout(0.2)
try:
    s.connect(('169.254.169.254', 80))
    raise AssertionError('network connected')
except OSError:
    pass
print('isolated')
'''
        with tempfile.TemporaryDirectory() as tmp:
            result = compute.run_execution(os.environ['JARVIS_TEST_VM_CONFIG'], Path(tmp),
                                           request(code=code), timeout=30)
            self.assertTrue(result['ok'], result)
            self.assertEqual(result['stdout'].strip(), 'isolated')


class GuestExportValidationTests(unittest.TestCase):
    def guest(self):
        spec = importlib.util.spec_from_file_location('guest_export', Path(__file__).parents[1] / 'code-mode-compute-guest.py')
        guest = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(guest)
        return guest

    def test_metadata_changes_after_first_stat_are_refused(self):
        from unittest.mock import patch
        import time
        guest = self.guest()
        for mode in ('hardlink', 'restored_mtime', 'replacement'):
            with self.subTest(mode=mode), tempfile.TemporaryDirectory() as tmp:
                root = Path(tmp)
                path = root / 'bad'; path.write_bytes(b'original')
                identity = path.stat()
                original_open, original_stat = os.open, os.fstat
                injected = [False]
                def redirect(name, *args, **kwargs):
                    return original_open(root if name == '/outputs' else name, *args, **kwargs)
                def race(fd):
                    info = original_stat(fd)
                    if not injected[0] and (info.st_dev, info.st_ino) == (identity.st_dev, identity.st_ino):
                        injected[0] = True
                        if mode == 'hardlink': os.link(path, root / 'alias')
                        elif mode == 'restored_mtime':
                            time.sleep(.01)
                            path.write_bytes(b'tampered')
                            os.utime(path, ns=(info.st_atime_ns, info.st_mtime_ns))
                        else:
                            path.rename(root / 'old'); path.write_bytes(b'tampered')
                    return info
                with patch.object(os, 'open', redirect), patch.object(os, 'fstat', race):
                    with self.assertRaisesRegex(ValueError, 'output_denied'):
                        guest.export_files(['bad'])
                self.assertTrue(injected[0])

    def test_actual_character_device_is_refused_without_reading(self):
        from unittest.mock import patch
        guest = self.guest()
        original_open = os.open
        def redirect(name, *args, **kwargs):
            return original_open('/dev' if name == '/outputs' else name, *args, **kwargs)
        with patch.object(os, 'open', redirect), self.assertRaisesRegex(ValueError, 'output_denied'):
            guest.export_files(['null'])


class WorkloadExitTests(unittest.TestCase):
    def test_resource_exceptions_are_distinct_from_ordinary_failures(self):
        import subprocess
        import sys
        guest_spec = importlib.util.spec_from_file_location('compute_guest', Path(__file__).parents[1] / 'code-mode-compute-guest.py')
        guest = importlib.util.module_from_spec(guest_spec)
        guest_spec.loader.exec_module(guest)
        self.assertTrue(hasattr(guest, 'WORKER_BOOTSTRAP'), 'worker needs a typed resource-exhaustion exit protocol')
        with tempfile.TemporaryDirectory() as tmp:
            wrapper = Path(tmp) / 'worker.py'
            wrapper.write_text(guest.WORKER_BOOTSTRAP)
            program = Path(tmp) / 'program.py'
            cases = [('raise MemoryError()', 125)]
            cases += [(f'import errno;raise OSError(errno.{name}, "fixture")', 125)
                      for name in ('ENOSPC', 'EDQUOT', 'EFBIG', 'ENOMEM', 'EAGAIN')]
            cases += [('raise PermissionError("fixture")', 1), ('raise RuntimeError("fixture")', 1),
                      ('raise SystemExit(7)', 7), ('print("done")', 0)]
            for source, expected in cases:
                with self.subTest(source=source):
                    program.write_text(source)
                    result = subprocess.run([sys.executable, '-I', '-B', wrapper, program],
                                            capture_output=True, timeout=5)
                    self.assertEqual(result.returncode, expected, result.stderr)


class ArtifactCapabilityTests(unittest.TestCase):
    def test_input_is_a_private_snapshot_and_id_is_task_scoped(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / 'one').mkdir(mode=0o700); (root / 'two').mkdir(mode=0o700)
            source = root / 'source.csv'; source.write_bytes(b'a,b\n1,2\n')
            one = compute.ArtifactStore(root / 'one')
            artifact = one.import_file(source, 'sheet.csv')
            self.assertNotIn(str(source), artifact['id'])
            source.write_bytes(b'changed')
            resolved = one.resolve_inputs([{'artifactId': artifact['id'], 'name': 'renamed.csv'}])
            self.assertEqual(resolved, {'renamed.csv': b'a,b\n1,2\n'})
            two = compute.ArtifactStore(root / 'two')
            with self.assertRaises(compute.ComputeError) as failure:
                two.resolve_inputs([{'artifactId': artifact['id'], 'name': 'stolen.csv'}])
            self.assertEqual(failure.exception.code, 'input_denied')

    def test_import_rejects_links_and_special_files_without_reading(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); (root / 'store').mkdir(mode=0o700)
            store = compute.ArtifactStore(root / 'store')
            source = root / 'source'; source.write_bytes(b'canary')
            link = root / 'link'; link.symlink_to(source)
            parent = root / 'parent'; parent.symlink_to(root, target_is_directory=True)
            fifo = root / 'fifo'; os.mkfifo(fifo)
            hardlink = root / 'hardlink'; os.link(source, hardlink)
            for path in (link, parent / 'source', fifo, source, hardlink, root):
                with self.subTest(path=path), self.assertRaises(compute.ComputeError):
                    store.import_file(path, 'input')

    def test_unknown_handle_cannot_be_interpreted_as_a_path(self):
        with tempfile.TemporaryDirectory() as tmp:
            store = compute.ArtifactStore(Path(tmp))
            for identifier in ('/etc/passwd', '../another-task/file', 'made-up-id'):
                with self.subTest(identifier=identifier), self.assertRaises(compute.ComputeError) as failure:
                    store.resolve_inputs([{'artifactId': identifier, 'name': 'file'}])
                self.assertEqual(failure.exception.code, 'input_denied')

    def test_untrusted_store_directory_is_refused(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); child = root / 'child'; child.mkdir(mode=0o755)
            link = root / 'link'; link.symlink_to(root, target_is_directory=True)
            for directory in (child, link):
                with self.subTest(directory=directory), self.assertRaises(compute.ComputeError):
                    compute.ArtifactStore(directory)


@unittest.skipUnless(os.environ.get('JARVIS_TEST_VM_CONFIG') and os.environ.get('JARVIS_TEST_COMPUTE_PIP'),
                     'requires KVM and a private pinned pip runtime')
class DependencyPreparationTests(unittest.TestCase):
    def test_task_reuses_real_environment_and_chains_an_exported_input(self):
        import json
        with tempfile.TemporaryDirectory() as tmp:
            backend = compute.VMBackend(os.environ['JARVIS_TEST_VM_CONFIG'], '/mnt/build/codex-vm',
                json.loads(Path(os.environ['JARVIS_TEST_COMPUTE_PIP']).read_text()))
            self.addCleanup(backend.close)
            artifacts = compute.ArtifactStore(Path(tmp))
            task = compute.ComputeTask(backend, artifacts, enabled=True)
            first = task.execute(request(dependencies=['openpyxl==3.1.5'], outputs=['input.xlsx'], code="""
import openpyxl
book = openpyxl.Workbook()
for n in (10,20,30): book.active.append([n])
book.save('/outputs/input.xlsx')
"""))
            self.assertTrue(first['ok'], first)
            second = task.execute(request(dependencies=['openpyxl==3.1.5'], outputs=['summary.json'],
                inputs=[{'artifactId': first['artifacts'][0]['id'], 'name': 'input.xlsx'}], code="""
import json,openpyxl
from pathlib import Path
book = openpyxl.load_workbook('/inputs/input.xlsx')
rows = list(book.active.values)
Path('/outputs/summary.json').write_text(json.dumps({'count':len(rows),'total':sum(row[0] for row in rows)}))
try:
    Path('/inputs/input.xlsx').write_bytes(b'mutated')
    raise AssertionError('input writable')
except OSError:
    pass
"""))
            self.assertTrue(second['ok'], second)
            self.assertTrue(second['environmentReused'])
            self.assertEqual(first['dependencyLock'], second['dependencyLock'])
            self.assertEqual(task.records[-1]['downloads']['requests'], 0)
            data = artifacts.resolve_inputs([{'artifactId': second['artifacts'][0]['id'], 'name': 'summary'}])
            self.assertEqual(json.loads(data['summary']), {'count':3,'total':60})

    def test_real_wheels_install_and_are_readonly_without_gateway_during_execution(self):
        import json
        bridge_spec = importlib.util.spec_from_file_location('compute_bridge', Path(__file__).parents[1] / 'codex-tool-bridge.py')
        bridge = importlib.util.module_from_spec(bridge_spec); bridge_spec.loader.exec_module(bridge)
        scratch = bridge.BuildScratch('/mnt/build/codex-vm')
        scratch.open()
        self.addCleanup(scratch.close)
        pip_runtime = json.loads(Path(os.environ['JARVIS_TEST_COMPUTE_PIP']).read_text())
        environment_id = 'a' * 32
        prepared = compute.run_preparation(os.environ['JARVIS_TEST_VM_CONFIG'], scratch.tmp, scratch.cache,
            pip_runtime, ['openpyxl==3.1.5'], environment_id, timeout=90)
        self.assertTrue(prepared['ok'], prepared)
        self.assertEqual({item['name'] for item in prepared['dependencyLock']}, {'openpyxl', 'et-xmlfile'})
        self.assertGreater(prepared['downloads']['requests'], 0)
        code = """
import openpyxl
from pathlib import Path
book = openpyxl.Workbook()
book.active.append([10]); book.active.append([20]); book.active.append([30])
book.save('/outputs/input.xlsx')
assert sum(row[0] for row in book.active.values) == 60
try:
    Path(openpyxl.__file__).write_text('mutated')
    raise AssertionError('dependency environment is writable')
except OSError:
    pass
try:
    assert not Path('/root/control').exists()
except PermissionError:
    pass
print('60')
"""
        executed = compute.run_execution(os.environ['JARVIS_TEST_VM_CONFIG'], scratch.tmp,
            request(code=code, outputs=['input.xlsx']), timeout=30,
            cache=scratch.cache, environment_id=environment_id)
        self.assertTrue(executed['ok'], executed)
        self.assertEqual(executed['stdout'].strip(), '60')
        self.assertTrue(executed['files']['input.xlsx'].startswith(b'PK'))
        self.assertEqual(executed['downloads']['requests'], 0)


class TaskLifecycleContractTests(unittest.TestCase):
    class Backend:
        fingerprint = 'trusted-runtime-fixture'
        def __init__(self):
            self.prepares = []; self.executions = []; self.preparation_failure = False
        def admit(self, deadline=None):
            from contextlib import nullcontext
            return nullcontext()
        def prepare(self, requirements, deadline):
            self.prepares.append((requirements, deadline))
            if self.preparation_failure:
                raise compute.ComputeError('dependency_unavailable', 'Unavailable fixture.')
            return {'environmentId': 'a' * 32,
                    'dependencyLock': [{'name': 'example', 'version': '1', 'sha256': 'a' * 64}],
                    'downloads': {'requests': 2, 'bytes': 64}}
        def execute(self, request, inputs, environment_id, deadline):
            self.executions.append((request, inputs, environment_id, deadline))
            return {'ok': True, 'runner': 'vm', 'exitCode': 0, 'error': None,
                    'stdout': 'computed', 'stderr': '', 'files': {'result.json': b'{"total":60}'},
                    'cleanupVerified': True}

    def setup_task(self, root, **options):
        backend = self.Backend()
        artifacts = compute.ArtifactStore(root)
        return backend, compute.ComputeTask(backend, artifacts, enabled=True, **options)

    def test_expiry_during_export_rolls_back_artifacts_before_returning(self):
        from unittest.mock import patch
        with tempfile.TemporaryDirectory() as tmp:
            now = [10.0]
            backend, task = self.setup_task(Path(tmp), clock=lambda: now[0])
            store = task.artifacts._store
            def expire_after_write(data, name):
                entry = store(data, name)
                now[0] = 13.0
                return entry
            with patch.object(task.artifacts, '_store', side_effect=expire_after_write):
                result = task.execute(request(outputs=['result.json']), host_deadline=12.0)
            self.assertFalse(result['ok'], result)
            self.assertEqual(result['error']['code'], 'timeout')
            self.assertEqual(result['artifacts'], [])
            self.assertEqual(task.artifacts.entries, {})
            self.assertFalse(any(len(path.name) == 32 for path in Path(tmp).iterdir()), 'expired export left a capability file')

    def test_invalid_private_deadlines_fail_before_backend_access(self):
        for deadline in (True, '12', float('nan'), float('inf'), 0, -1):
            with self.subTest(deadline=deadline), tempfile.TemporaryDirectory() as tmp:
                backend, task = self.setup_task(Path(tmp))
                with self.assertRaises(compute.ComputeError) as failure:
                    task.execute(request(), host_deadline=deadline)
                self.assertEqual(failure.exception.code, 'bad_args')
                self.assertEqual(backend.executions, [])
                self.assertEqual(backend.prepares, [])
                self.assertEqual(list(Path(tmp).iterdir()), [])

    def test_host_deadline_cannot_reset_or_extend_task_or_call_budget(self):
        with tempfile.TemporaryDirectory() as tmp:
            now = [10.0]
            backend, task = self.setup_task(Path(tmp), call_timeout=3, task_timeout=30,
                                           host_deadline=35.0, clock=lambda: now[0])
            self.assertEqual(task.deadline, 35.0)
            result = task.execute(request(dependencies=['example'], outputs=['result.json']), host_deadline=12.0)
            self.assertTrue(result['ok'], result)
            self.assertEqual(backend.prepares[0][1], 12.0)
            self.assertEqual(backend.executions[0][-1], 12.0)
            self.assertEqual(task.records[0].get('deadlineMonotonic'), 12.0)
            self.assertEqual(task.records[0].get('taskDeadlineMonotonic'), 35.0)
            now[0] = 20.0
            self.assertTrue(task.execute(request(outputs=['result.json']), host_deadline=1000.0)['ok'])
            self.assertEqual(backend.executions[-1][-1], 23.0)
            now[0] = 34.0
            self.assertTrue(task.execute(request(outputs=['result.json']), host_deadline=1000.0)['ok'])
            self.assertEqual(backend.executions[-1][-1], 35.0)
            count = len(backend.executions)
            expired = task.execute(request(), host_deadline=33.0)
            self.assertEqual(expired['error']['code'], 'timeout')
            self.assertEqual(expired['runner'], 'none')
            self.assertEqual(len(backend.executions), count)

    def test_exhausted_shared_storage_has_typed_resource_limit(self):
        from unittest.mock import patch
        with tempfile.TemporaryDirectory() as tmp:
            backend = compute.VMBackend('/missing-runtime', Path(tmp))
            self.addCleanup(backend.close)
            with patch.object(backend.scratch, 'open', side_effect=backend.bridge.Readiness('build_scratch_space', tmp)):
                with self.assertRaises(compute.ComputeError) as failure:
                    with backend.admit():
                        self.fail('exhausted admission must not start work')
            self.assertEqual(failure.exception.code, 'resource_limit')

    def test_call_limit_also_bounds_retained_audit_records(self):
        with tempfile.TemporaryDirectory() as tmp:
            backend, task = self.setup_task(Path(tmp))
            for _ in range(25):
                self.assertTrue(task.execute(request(outputs=['result.json']))['ok'])
            with self.assertRaises(compute.ComputeError) as failure:
                task.execute(request(outputs=['result.json']))
            self.assertEqual(failure.exception.code, 'resource_limit')
            self.assertEqual(len(task.records), 25)
            self.assertEqual(len(backend.executions), 25)

    def test_policy_fingerprint_and_phase_timings_follow_host_policy_and_clock(self):
        import hashlib
        import json
        with tempfile.TemporaryDirectory() as tmp:
            now = [10.0]
            backend, task = self.setup_task(Path(tmp), clock=lambda: now[0], task_timeout=60)
            prepare, execute, publish = backend.prepare, backend.execute, task.artifacts.publish
            def timed_prepare(*args):
                now[0] += 2
                return prepare(*args)
            def timed_execute(*args):
                now[0] += 3
                return execute(*args)
            def timed_publish(*args, **kwargs):
                now[0] += 4
                return publish(*args, **kwargs)
            backend.prepare, backend.execute, task.artifacts.publish = timed_prepare, timed_execute, timed_publish
            self.assertTrue(task.execute(request(dependencies=['example'], outputs=['result.json']))['ok'])
            record = task.records[-1]
            self.assertIn('policyFingerprint', record, 'audit cannot identify the enforced policy')
            policy = record['policy']
            expected = hashlib.sha256(json.dumps(policy, sort_keys=True, separators=(',', ':')).encode()).hexdigest()
            self.assertEqual(record['policyFingerprint'], expected)
            self.assertEqual(policy['taskTimeoutSecs'], 60)
            self.assertEqual(policy['callTimeoutSecs'], 600)
            self.assertEqual([(phase['phase'], phase['elapsedSecs']) for phase in record['phases']],
                             [('admission', 0), ('prepare', 2), ('execute', 3), ('export', 4)])
            self.assertEqual(record['elapsedSecs'], 9)
            previous = record['startedMonotonic']
            for phase in record['phases']:
                self.assertEqual(phase['startedMonotonic'], previous)
                previous = phase['endedMonotonic']
            self.assertEqual(previous - record['startedMonotonic'], record['elapsedSecs'])
            first = record['policyFingerprint']
            self.assertTrue(task.execute(request(dependencies=['example'], outputs=['result.json']))['ok'])
            self.assertEqual(task.records[-1]['policyFingerprint'], first)
            self.assertNotIn('prepare', [phase['phase'] for phase in task.records[-1]['phases']])
        with tempfile.TemporaryDirectory() as tmp:
            _, other = self.setup_task(Path(tmp), task_timeout=61)
            other.execute(request(outputs=['result.json']))
            self.assertNotEqual(other.records[-1]['policyFingerprint'], first)

    def test_preparation_audit_preserves_private_diagnostics_on_success_and_failure(self):
        import json
        for failed in (False, True):
            with self.subTest(failed=failed), tempfile.TemporaryDirectory() as tmp:
                root = Path(tmp)
                backend, task = self.setup_task(root)
                prepare = backend.prepare
                outcome = {'runner': 'vm', 'cleanupVerified': True,
                           'downloads': {'requests': 3, 'bytes': 123},
                           'error': 'dependency_unavailable' if failed else None,
                           '_logs': {'stdout': b'private installer output', 'stderr': b'private preparation canary'}}
                def with_diagnostics(*args):
                    if failed:
                        error = compute.ComputeError('dependency_unavailable', 'fixed diagnostic', runner='vm')
                        error.preparation = outcome
                        raise error
                    result = prepare(*args)
                    result['downloads'] = outcome['downloads']
                    result['preparation'] = outcome
                    return result
                backend.prepare = with_diagnostics
                result = task.execute(request(dependencies=['example'], outputs=['result.json']))
                self.assertEqual(result['ok'], not failed)
                record = task.records[-1]
                metadata = record.get('preparation')
                self.assertIsNotNone(metadata, 'completed preparation diagnostics were discarded')
                self.assertEqual(record['downloads'], outcome['downloads'])
                self.assertEqual(metadata['downloads'], outcome['downloads'])
                self.assertTrue(metadata['cleanupVerified'])
                self.assertGreaterEqual(metadata['elapsedSecs'], 0)
                self.assertEqual(metadata['error'], outcome['error'])
                self.assertNotIn('private preparation canary', json.dumps(result))
                self.assertNotIn('private preparation canary', (root / 'audit.json').read_text())
                for stream, raw in outcome['_logs'].items():
                    path = root / metadata['logs'][stream]['file']
                    self.assertEqual(path.read_bytes(), raw)
                    self.assertEqual(path.stat().st_mode & 0o777, 0o600)
                if not failed:
                    self.assertTrue(task.execute(request(dependencies=['example'], outputs=['result.json']))['environmentReused'])
                    self.assertIsNone(task.records[-1]['preparation'])
                    self.assertEqual(task.records[-1]['downloads'], {'requests': 0, 'bytes': 0})
                    self.assertTrue(all('preparation' not in environment for environment in task.environments.values()))

    def test_private_audit_preserves_full_logs_without_source_in_metadata(self):
        import json
        import stat
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            backend, task = self.setup_task(root)
            original = backend.execute
            def execute(*args):
                result = original(*args)
                result['stdout'] = 'x' * 70000
                result['stderr'] = 'private-log-canary'
                return result
            backend.execute = execute
            result = task.execute(request(code='# source-only-canary', outputs=['result.json']))
            self.assertTrue(result['ok'])
            self.assertEqual(len(result['stdout']), 65536)
            audit = root / 'audit.json'
            self.assertTrue(audit.exists(), 'execution metadata must survive helper exit')
            metadata = json.loads(audit.read_text())
            self.assertNotIn('source-only-canary', audit.read_text())
            self.assertNotIn('private-log-canary', audit.read_text())
            record = metadata['records'][0]
            self.assertEqual(record['logs']['stdout']['bytes'], 70000)
            self.assertTrue(record['logs']['stdout']['responseTruncated'])
            for name, expected in [('stdout', b'x' * 70000), ('stderr', b'private-log-canary')]:
                path = root / record['logs'][name]['file']
                self.assertEqual(path.read_bytes(), expected)
                self.assertEqual(stat.S_IMODE(path.stat().st_mode), 0o600)
            self.assertEqual(stat.S_IMODE(audit.stat().st_mode), 0o600)

    def test_interrupted_execution_is_recorded_and_cleanup_stays_unverified(self):
        import json
        class Interrupted(BaseException):
            pass
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            backend, task = self.setup_task(root)
            def execute(*args):
                audit = root / 'audit.json'
                self.assertTrue(audit.exists(), 'write active lease before running the guest')
                self.assertEqual(json.loads(audit.read_text())['active']['phase'], 'execute')
                raise Interrupted()
            backend.execute = execute
            with self.assertRaises(Interrupted):
                task.execute(request())
            self.assertEqual(len(task.records), 1, 'cancellation lost its execution record')
            record = task.records[0]
            self.assertEqual(record['error']['code'], 'cancelled')
            self.assertFalse(record['cleanupVerified'])
            audit = json.loads((root / 'audit.json').read_text())
            self.assertIsNone(audit['active'])
            self.assertEqual(audit['records'], task.records)

    def test_identical_constraints_reuse_lock_and_changed_constraints_prepare_again(self):
        with tempfile.TemporaryDirectory() as tmp:
            backend, task = self.setup_task(Path(tmp))
            spec = request(dependencies=['Example==1'], outputs=['result.json'])
            first = task.execute(spec); second = task.execute(spec)
            self.assertTrue(first['ok']); self.assertTrue(second['ok'])
            self.assertFalse(first['environmentReused']); self.assertTrue(second['environmentReused'])
            self.assertEqual(first['dependencyLock'], second['dependencyLock'])
            self.assertEqual(len(backend.prepares), 1)
            self.assertEqual(len(backend.executions), 2)
            task.execute(request(dependencies=['Example>=1'], outputs=['result.json']))
            self.assertEqual(len(backend.prepares), 2)

    def test_output_handle_can_be_selected_as_next_input(self):
        with tempfile.TemporaryDirectory() as tmp:
            backend, task = self.setup_task(Path(tmp))
            first = task.execute(request(outputs=['result.json']))
            self.assertTrue(first['ok'])
            artifact = first['artifacts'][0]
            task.execute(request(inputs=[{'artifactId': artifact['id'], 'name': 'prior.json'}], outputs=['result.json']))
            self.assertEqual(backend.executions[-1][1], {'prior.json': b'{"total":60}'})

    def test_missing_input_fails_before_backend_or_resolution(self):
        with tempfile.TemporaryDirectory() as tmp:
            backend, task = self.setup_task(Path(tmp))
            result = task.execute(request(inputs=[{'artifactId': 'foreign', 'name': 'file'}]))
            self.assertFalse(result['ok']); self.assertEqual(result['error']['code'], 'input_denied')
            self.assertEqual(backend.prepares, []); self.assertEqual(backend.executions, [])

    def test_failed_preparation_is_never_reused(self):
        with tempfile.TemporaryDirectory() as tmp:
            backend, task = self.setup_task(Path(tmp))
            backend.preparation_failure = True
            first = task.execute(request(dependencies=['example'], outputs=['result.json']))
            self.assertFalse(first['ok'])
            backend.preparation_failure = False
            second = task.execute(request(dependencies=['example'], outputs=['result.json']))
            self.assertTrue(second['ok']); self.assertFalse(second['environmentReused'])
            self.assertEqual(len(backend.prepares), 2)

    def test_disabled_and_expired_tasks_never_start_backend(self):
        with tempfile.TemporaryDirectory() as tmp:
            now = [10.0]
            backend, task = self.setup_task(Path(tmp), task_timeout=2, clock=lambda: now[0])
            now[0] = 12.0
            result = task.execute(request())
            self.assertEqual(result['error']['code'], 'timeout')
            disabled = compute.ComputeTask(backend, task.artifacts, enabled=False)
            self.assertEqual(disabled.execute(request())['error']['code'], 'compute_disabled')
            self.assertEqual(backend.prepares, []); self.assertEqual(backend.executions, [])


class HostProtocolTests(unittest.TestCase):
    def test_disabled_helper_rejects_compute_without_reading_runtime(self):
        import json
        import subprocess
        import sys
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); (root / 'artifacts').mkdir(mode=0o700)
            policy = root / 'policy.json'
            policy.write_text(json.dumps({'runtime': '/absent/runtime.json', 'scratch': '/absent/scratch',
                'artifactRoot': str(root / 'artifacts'), 'enabled': False, 'callTimeoutSecs': 600,
                'taskTimeoutSecs': 1800, 'inputFiles': {}})); policy.chmod(0o600)
            result = subprocess.run([sys.executable, '-I', str(Path(compute.__file__).resolve()), '--serve', str(policy)],
                input=json.dumps({'execute':request()})+'\n'+json.dumps({'close':True})+'\n',
                text=True, capture_output=True, timeout=5)
            self.assertEqual(result.returncode, 0, result.stderr)
            frames = [json.loads(line) for line in result.stdout.splitlines()]
            self.assertTrue(frames[0]['ready'])
            self.assertEqual(frames[1]['result']['error']['code'], 'compute_disabled')
            self.assertTrue(frames[2]['cleanupVerified'])


if __name__ == '__main__':
    unittest.main()
